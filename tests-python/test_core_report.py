from pathlib import Path
from tempfile import TemporaryDirectory
import json
import unittest
from openpyxl import load_workbook
from qt_research.core_report import overview, episodes, write_report
from qt_research.event_engine import FixedQuantityStrategy, run_events
from qt_research.ledger import AccountLedger, Fill
from qt_research.execution import Bar
from qt_research.orders import SecurityRule
from qt_research.local_cli import inspect_result


class ReportTests(unittest.TestCase):
    def test_first_fee_and_year_boundary_and_relative_excess(self):
        rows = [{"time": t, "equity": v, "baseline": base} for t, v, base in
                (("2024-12-30T15:00:00+08:00", 1000, True),
                 ("2024-12-31T15:00:00+08:00", 995, False),
                 ("2025-01-02T15:00:00+08:00", 1094.5, False))]
        metrics, daily, yearly = overview(rows, 1000, [], {"2024-12-31": 1, "2025-01-02": 1.05})
        self.assertAlmostEqual(metrics["strategy_return"], .0945)
        self.assertAlmostEqual(yearly["2024"], -.005)
        self.assertAlmostEqual(yearly["2025"], .1)
        self.assertAlmostEqual(metrics["excess_return"], 1.0945 / 1.05 - 1)
        self.assertAlmostEqual(metrics["max_drawdown"], .005)
        self.assertIsNone(metrics["trade_win_rate"])

    def test_export_workbook_checksums_and_preserve_existing(self):
        bars = [Bar(f"2024-01-02T14:0{i}:00+08:00", "ETF", 10, 400) for i in range(3)]
        result = run_events(bars, FixedQuantityStrategy("ETF", 100), {"ETF": SecurityRule()}, 2000)
        with TemporaryDirectory() as temp:
            out = Path(temp) / "run"
            write_report(out, result, {"description": "Test"})
            verification = inspect_result(out)
            self.assertEqual(verification["summary"]["fills"], 1)
            wb = load_workbook(out / "transactions.xlsx", read_only=True)
            self.assertEqual(wb.sheetnames, ["委托", "成交", "已平仓交易", "账户", "持仓", "收益概况"])
            self.assertEqual(wb["成交"].max_row, 2)
            self.assertEqual(wb["账户"].max_row, 2)
            columns = [cell.value for cell in wb["成交"][1]]
            self.assertIsInstance(wb["成交"].cell(2, columns.index("price") + 1).value, (int, float))
            wb.close()
            saved = (out / "summary.json").read_bytes()
            with self.assertRaises(FileExistsError):
                write_report(out, result, {})
            self.assertEqual((out / "summary.json").read_bytes(), saved)
            with (out / "equity.csv").open("a") as stream:
                stream.write("tampered\n")
            with self.assertRaises(ValueError):
                inspect_result(out)

    def test_closed_position_groups_partial_fills_as_one_trade(self):
        ledger = AccountLedger(5000)
        for fid, side, qty, price in (("b", "buy", 200, 10), ("s1", "sell", 50, 11), ("s2", "sell", 150, 12)):
            ledger.apply_fill(Fill(fid, fid, "2024-01-02", "ETF", side, qty, price, 5))
        closed = episodes(ledger.entries)
        self.assertEqual(len(closed), 1)
        self.assertEqual(closed[0]["fills"], 3)
        self.assertEqual(closed[0]["pnl"], 335)

    def test_zero_volatility_and_missing_benchmark_are_unavailable(self):
        rows = [{"time": "2024-01-02T15:00:00+08:00", "equity": 1000, "baseline": False}]
        metrics, _, _ = overview(rows, 1000, [])
        self.assertIsNone(metrics["sharpe"])
        self.assertIsNone(metrics["beta"])
        self.assertIsNone(metrics["trade_win_rate"])


if __name__ == "__main__":
    unittest.main()
