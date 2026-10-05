"""Read-only presentation of verified fill ledgers and JoinQuant orders."""
import csv
from decimal import Decimal
from datetime import date


def number(value):
    return Decimal(str(value)) if value not in (None, '') else None


def table(path, sheet):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            return []
        it = wb[sheet].iter_rows(values_only=True)
        headers = next(it, [])
        return [dict(zip(headers, row)) for row in it]
    finally:
        wb.close()


def names(cat, run, verified):
    result = {}
    candidates = [run] + [r for r in cat.rows('runs') if r['strategy_id'] == run['strategy_id'] and r['platform'] == 'joinquant' and r['id'] != run['id']]
    for r in candidates:
        for a in cat.detail(r['id'])['artifacts']:
            if a['name'].endswith('transactions.normalized.csv'):
                with verified(cat, r['id'], a['id']).open() as f:
                    for x in csv.DictReader(f):
                        if x.get('security_name') and x.get('code'):
                            result.setdefault(x['code'].split('.')[0], x['security_name'])
    return result


def trade_rows(cat, rid, verified, group='order'):
    if group not in ('order', 'fill'):
        raise ValueError('Unknown trade view')
    run = cat.detail(rid)
    arts = run['artifacts']
    excel = next((a for a in arts if a['name'] == 'transactions.xlsx'), None)
    rows = []
    if excel:
        path = verified(cat, rid, excel['id'])
        fills = table(path, '成交')
        orders = {x['order_id']: x for x in table(path, '委托')}
        instruments = names(cat, run, verified)
        groups = {}
        for i, fill in enumerate(fills):
            key = str(fill['order_id']) if group == 'order' else str(i)
            groups.setdefault(key, []).append(fill)
        for key, fs in groups.items():
            fs.sort(key=lambda x: str(x['time']))
            first = fs[0]
            order = orders.get(first['order_id'], {})
            quantity = sum((abs(number(x['quantity'])) for x in fs), Decimal(0))
            gross = sum((abs(number(x['quantity'])) * number(x['price']) for x in fs), Decimal(0))
            fees = sum((number(x['fee']) or Decimal(0) for x in fs), Decimal(0))
            pnl = sum((number(x.get('realized_pnl')) or Decimal(0) for x in fs), Decimal(0))
            sign = -1 if first['side'] == 'sell' else 1
            ordered = number(order.get('quantity'))
            filled = number(order.get('filled'))
            state = order.get('status', '')
            status = '全部成交' if state == 'filled' else '部分成交' if state in ('open', 'accepted', 'partially_filled') else '部分成交后取消' if state == 'cancelled' and filled else state or '未记录'
            rows.append(dict(date=str(first['time'])[:10], time=str(first['time'])[11:19],
                symbol=first['symbol'], security_name=instruments.get(first['symbol'].split('.')[0], ''),
                side='sell' if sign == -1 else 'buy', order_type='限价单' if order.get('limit_price') is not None else None,
                quantity=float(quantity * sign), price=float(gross / quantity) if quantity else None,
                value=float(gross * sign), pnl=float(pnl), fee=float(fees),
                ordered_quantity=float(ordered * sign) if ordered is not None else None,
                order_price=order.get('limit_price'), status=status, last_update=str(fs[-1]['time']),
                fill_count=len(fs), order_id=first['order_id'], raw=dict(order=order, fills=fs)))
    else:
        a = next((a for a in arts if a['name'].endswith('transactions.normalized.csv')), None)
        if not a:
            raise KeyError('Transactions unavailable')
        with verified(cat, rid, a['id']).open() as f:
            for i, x in enumerate(csv.DictReader(f)):
                n = lambda key: float(number(x.get(key))) if number(x.get(key)) is not None else None
                side = 'sell' if x.get('side') in ('卖', 'sell') else 'buy' if x.get('side') in ('买', 'buy') else x.get('side')
                rows.append(dict(date=x.get('date'), time=x.get('order_time'), symbol=x.get('code'),
                    security_name=x.get('security_name'), side=side, order_type=x.get('order_type'),
                    quantity=n('filled_amount'), price=n('filled_price'), value=n('filled_value'),
                    pnl=n('close_pnl'), fee=n('fee'), ordered_quantity=n('ordered_amount'),
                    order_price=n('order_price'), status=x.get('status'), last_update=x.get('last_update'),
                    fill_count=None, order_id=None, raw=x))
    return rows, run['platform']


def select(rows, start='', end='', side='', keyword='', offset=0, limit=50):
    for day in (start, end):
        if day:
            date.fromisoformat(day)
    if start and end and start > end:
        raise ValueError('Start date must not exceed end date')
    if side not in ('', 'buy', 'sell'):
        raise ValueError('Unknown trade side')
    needle = keyword.casefold()
    filtered = [r for r in rows if (not start or r['date'] >= start) and
        (not end or r['date'] <= end) and (not side or r['side'] == side) and
        (not needle or needle in (str(r['symbol']) + ' ' + str(r.get('security_name') or '')).casefold())]
    filtered.sort(key=lambda r: (r['date'] or '', r['time'] or ''))
    return dict(rows=filtered[offset:offset + limit], total=len(filtered), all_total=len(rows))
