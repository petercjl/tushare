"""Loopback-only catalog viewer; no execution or trading endpoints."""
import argparse
import csv
import hashlib
import json
import mimetypes
import os
from pathlib import Path
import signal
import subprocess
import sys
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlsplit,parse_qs
from urllib.request import urlopen
from .research_catalog import Catalog,ROOT,digest

PORT=8766

def artifact(cat,rid,aid):
    row=next((a for a in cat.detail(rid)['artifacts'] if a['id']==aid),None)
    if row is None:raise KeyError('Artifact not found')
    p=Path(row['path'])
    if digest(p)!=row['sha256']:raise ValueError('Artifact checksum mismatch')
    return p

def curve(cat,rid):
    run=cat.detail(rid);arts=run['artifacts']
    e=next((a for a in arts if a['name']=='equity.csv'),None)
    if e:
        p=artifact(cat,rid,e['id'])
        with p.open() as f:rows=list(csv.DictReader(f))
        def percent(row,key):
            value=row.get(key)
            return float(value)*100 if value not in (None,'') else None
        return [dict(date=r.get('date') or r['time'][:10],equity=float(r['equity']),
                     value=float(r['strategy_return'])*100,
                     benchmark_value=percent(r,'benchmark_return'),
                     excess_value=percent(r,'excess_return')) for r in rows]
    e=next((a for a in arts if a['name'].endswith('logs.audit.jsonl')),None)
    if not e:return []
    # JoinQuant export percentages are already in percent units. Preserve the
    # strategy equity from its audit log and join only benchmark/excess by date.
    exported=next((a for a in arts if a['name'].endswith('result.normalized.csv')),None)
    lookup={}
    if exported:
        with artifact(cat,rid,exported['id']).open() as f:
            for item in csv.DictReader(f):
                lookup[item['date']]=item
    capital=run['params']['initial_cash'];rows=[]
    for line in artifact(cat,rid,e['id']).read_text().splitlines():
        x=json.loads(line)
        if x.get('event')=='daily_account':
            value=x.get('equity',x.get('portfolio_value'))
            if value is None:raise ValueError('Unknown JQ equity schema')
            day=x.get('dt',x.get('time'))[:10];item=lookup.get(day,{})
            def exported_percent(key):
                raw=item.get(key)
                return float(raw) if raw not in (None,'') else None
            rows.append(dict(date=day,equity=value,value=(value/capital-1)*100,
                             benchmark_value=exported_percent('benchmark_return_pct'),
                             excess_value=exported_percent('excess_return_pct')))
    return rows

def handler(catalog,port):
    class Handler(BaseHTTPRequestHandler):
        def log_message(self,*args):pass
        def send(self,body,kind='application/json; charset=utf-8',status=200,download=None):
            if not isinstance(body,bytes):body=body.encode()
            self.send_response(status);self.send_header('Content-Type',kind);self.send_header('Content-Length',str(len(body)))
            self.send_header('X-Content-Type-Options','nosniff');self.send_header('Cache-Control','no-store')
            if download:self.send_header('Content-Disposition','attachment; filename="'+download+'"')
            self.end_headers();self.wfile.write(body)
        def do_GET(self):
            if self.headers.get('Host') not in (f'127.0.0.1:{port}',f'localhost:{port}'):
                self.send('Invalid host',status=403);return
            u=urlsplit(self.path);q=parse_qs(u.query);get=lambda key:q[key][0]
            try:
                if u.path=='/':return self.send(Path(__file__).with_name('dashboard.html').read_bytes(),'text/html; charset=utf-8')
                if u.path=='/plotly.js':
                    from plotly.offline import get_plotlyjs
                    return self.send(get_plotlyjs(),'text/javascript; charset=utf-8')
                if u.path=='/api/health':value=dict(service='quant-research-dashboard',schema=1,pid=os.getpid(),database=str(catalog.path))
                elif u.path=='/api/catalog':value=dict(stats=catalog.stats(),strategies=catalog.rows('strategies'),versions=catalog.rows('versions'),runs=catalog.rows('runs'),comparisons=catalog.rows('comparisons'))
                elif u.path=='/api/run':value=catalog.detail(get('id'))
                elif u.path=='/api/curve':value=curve(catalog,get('id'))
                elif u.path=='/api/trades':
                    from .trade_view import trade_rows,select
                    group=q.get('group',['order'])[0]
                    rows,platform=trade_rows(catalog,get('id'),artifact,group)
                    value=select(rows,start=q.get('start',[''])[0],end=q.get('end',[''])[0],
                                 side=q.get('side',[''])[0],keyword=q.get('keyword',[''])[0],
                                 offset=max(0,int(q.get('offset',['0'])[0])),
                                 limit=min(200,max(1,int(q.get('limit',['50'])[0]))))
                    value.update(platform=platform,group=group)
                elif u.path=='/api/code':
                    row=next((v for v in catalog.rows('versions') if v['id']==get('id')),None)
                    if not row or not row['code_path']:raise KeyError('Code snapshot unavailable')
                    p=Path(row['code_path'])
                    if digest(p)!=row['code_hash']:raise ValueError('Code snapshot modified')
                    return self.send(p.read_bytes(),'text/plain; charset=utf-8')
                elif u.path=='/api/transactions':
                    run=catalog.detail(get('id'));a=next((a for a in run['artifacts'] if a['name']=='transactions.xlsx'),None)
                    offset=max(0,int(q.get('offset',['0'])[0]));limit=min(1000,max(1,int(q.get('limit',['200'])[0])))
                    if a:
                        import openpyxl
                        wb=openpyxl.load_workbook(artifact(catalog,run['id'],a['id']),read_only=True,data_only=True)
                        try:
                            sheet=get('sheet') if 'sheet' in q else '成交'
                            if sheet not in wb.sheetnames:raise KeyError('Sheet not found')
                            it=wb[sheet].iter_rows(values_only=True);headers=next(it);data=[]
                            for i,r in enumerate(it):
                                if offset<=i<offset+limit:data.append(list(r))
                            value=dict(headers=list(headers),rows=data,total=wb[sheet].max_row-1,sheets=wb.sheetnames)
                        finally:wb.close()
                    else:
                        a=next((a for a in run['artifacts'] if a['name'].endswith('transactions.normalized.csv')),None)
                        if not a:raise KeyError('Transactions unavailable')
                        with artifact(catalog,run['id'],a['id']).open() as f:data=list(csv.reader(f))
                        value=dict(headers=data[0],rows=data[1+offset:1+offset+limit],total=len(data)-1,sheets=[])
                elif u.path=='/artifact':
                    p=artifact(catalog,get('run'),get('id'));kind=mimetypes.guess_type(p.name)[0] or 'application/octet-stream'
                    return self.send(p.read_bytes(),kind,download=p.name if p.suffix not in ('.html','.txt','.log','.json','.csv','.py','.jsonl') else None)
                elif u.path=='/comparison':
                    row=next((r for r in catalog.rows('comparisons') if r['id']==get('id')),None)
                    if not row:raise KeyError('Comparison not found')
                    p=Path(row['out'])/'report.html';checks=json.loads((p.parent/'checksums.json').read_text())
                    if digest(p)!=checks['report.html']:raise ValueError('Comparison report modified')
                    return self.send(p.read_bytes(),'text/html; charset=utf-8')
                else:return self.send('Not found',status=404)
                self.send(json.dumps(value,ensure_ascii=False,allow_nan=False,default=str))
            except (KeyError,FileNotFoundError) as e:self.send(json.dumps(dict(error=str(e))),status=404)
            except Exception as e:self.send(json.dumps(dict(error=str(e))),status=400)
    return Handler

def serve(port=PORT,db=None):
    cat=Catalog(db) if db else Catalog()
    server=ThreadingHTTPServer(('127.0.0.1',port),handler(cat,port));server.daemon_threads=True
    server.serve_forever()

def status(port=PORT):
    try:
        with urlopen(f'http://127.0.0.1:{port}/api/health',timeout=2) as r:value=json.load(r)
        if value.get('service')!='quant-research-dashboard':raise ValueError('Unexpected service')
        return dict(running=True,url=f'http://127.0.0.1:{port}',**value)
    except Exception:return dict(running=False,url=f'http://127.0.0.1:{port}')

def main(argv=None):
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('action',choices=['start','serve','status','stop']);p.add_argument('--port',type=int,default=PORT)
    args=p.parse_args(argv)
    if not 1024<=args.port<=65535:raise ValueError('Port outside supported range')
    if args.action=='serve':return serve(args.port)
    result=status(args.port)
    if args.action=='status':return result
    if args.action=='stop':
        if result['running']:os.kill(result['pid'],signal.SIGTERM)
        return dict(stopped=result['running'])
    if result['running']:return result
    folder=ROOT/'state/services';folder.mkdir(parents=True,exist_ok=True,mode=0o700)
    log=folder/(str(time.time_ns())+'-dashboard.log')
    with log.open('xb') as f:
        os.chmod(log,0o600)
        child=subprocess.Popen([sys.executable,'-m','tushare_env.cli_service','--port',str(args.port)],stdin=subprocess.DEVNULL,stdout=f,stderr=f,start_new_session=True,cwd=ROOT)
    for _ in range(40):
        result=status(args.port)
        if result['running']:return dict(result,log=str(log))
        if child.poll() is not None:raise RuntimeError('Dashboard failed to start; see '+str(log))
        time.sleep(.1)
    raise RuntimeError('Dashboard start timed out; see '+str(log))
