"""Auditable reports for the quantity/event engine; independent of legacy metrics."""
from dataclasses import asdict
from pathlib import Path
import hashlib
import html
import json
import math
from decimal import Decimal

import numpy as np
import pandas as pd


def json_value(value):
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, dict):
        return {str(k): json_value(v) for k, v in value.items()}
    if isinstance(value, (tuple, list)):
        return [json_value(v) for v in value]
    if isinstance(value, (np.integer, np.floating)):
        return value.item()
    return value


def episodes(entries):
    """One closed trade = symbol position from zero through return to zero."""
    active, rows = {}, []
    for entry in entries:
        fill = entry.fill
        state = active.setdefault(fill.symbol, {"symbol": fill.symbol, "opened": fill.time,
                                                "buy_cost": Decimal("0"), "pnl": Decimal("0"),
                                                "fills": 0})
        state["fills"] += 1
        if fill.side == "buy":
            state["buy_cost"] += fill.price * fill.quantity + fill.fee
        else:
            state["pnl"] += entry.realized_pnl
        if entry.quantity_after == 0:
            rows.append(dict(state, closed=fill.time,
                             return_rate=float(state["pnl"] / state["buy_cost"])))
            del active[fill.symbol]
    return rows


def overview(equity, initial_cash, closed, benchmark=None, trading_days=252, risk_free=0.0):
    if trading_days <= 0 or risk_free <= -1 or not math.isfinite(risk_free):
        raise ValueError("Invalid performance configuration")
    actual = pd.DataFrame(equity)
    actual = actual.loc[~actual.baseline].copy()
    actual["date"] = pd.to_datetime(actual.time).dt.date.astype(str)
    daily = actual.groupby("date", sort=True).last()
    daily["equity"] = pd.to_numeric(daily.equity)
    initial = float(initial_cash)
    if daily.empty or initial <= 0 or (daily.equity <= 0).any():
        raise ValueError("Positive initial and daily equity required")
    wealth = daily.equity.to_numpy() / initial
    returns = np.diff(np.r_[1.0, wealth]) / np.r_[1.0, wealth[:-1]]
    rf = (1 + risk_free) ** (1 / trading_days) - 1
    daily["nav"] = wealth
    daily["daily_return"] = returns
    sigma = float(np.std(returns, ddof=1)) if len(returns) > 1 else 0.0
    downside = float(np.sqrt(np.mean(np.minimum(returns - rf, 0) ** 2)))
    path = np.r_[1.0, wealth]
    peak = np.maximum.accumulate(path)
    drawdowns = path / peak - 1
    trough = int(np.argmin(drawdowns))
    top = int(np.argmax(path[:trough + 1]))
    def label(index):
        return "initial_capital" if index == 0 else str(daily.index[index - 1])
    pnl = [float(row["pnl"]) for row in closed]
    profits, losses = [v for v in pnl if v > 0], [v for v in pnl if v < 0]
    metrics = {
        "strategy_return": float(wealth[-1] - 1),
        "annualized_return": float(wealth[-1] ** (trading_days / len(wealth)) - 1),
        "strategy_volatility": sigma * math.sqrt(trading_days),
        "sharpe": float(np.mean(returns - rf) / sigma * math.sqrt(trading_days)) if sigma else None,
        "sortino": float(np.mean(returns - rf) / downside * math.sqrt(trading_days)) if downside else None,
        "max_drawdown": float(-drawdowns[trough]),
        "max_drawdown_start": label(top) if trough else None,
        "max_drawdown_end": label(trough) if trough else None,
        "daily_win_rate": float(np.mean(returns > 0)),
        "trade_win_rate": len(profits) / len(pnl) if pnl else None,
        "profit_loss_ratio": float(np.mean(profits) / abs(np.mean(losses))) if profits and losses else None,
        "winning_trades": len(profits), "losing_trades": len(losses),
        "flat_trades": len(pnl) - len(profits) - len(losses), "closed_trades": len(pnl),
        "days": len(daily), "benchmark_return": None, "excess_return": None,
        "alpha": None, "beta": None, "benchmark_volatility": None,
        "information_ratio": None, "daily_excess_return": None,
        "excess_max_drawdown": None, "excess_sharpe": None,
    }
    daily["strategy_return"] = wealth - 1
    if benchmark is not None:
        bench = np.array([float(benchmark[str(day)]) for day in daily.index])
        if not np.isfinite(bench).all() or (bench <= 0).any():
            raise ValueError("Invalid benchmark normalized wealth")
        rb = np.diff(np.r_[1.0, bench]) / np.r_[1.0, bench[:-1]]
        relative = wealth / bench
        er = np.diff(np.r_[1.0, relative]) / np.r_[1.0, relative[:-1]]
        active = returns - rb
        sd_b = float(np.std(rb, ddof=1)) if len(rb) > 1 else 0.0
        sd_active = float(np.std(active, ddof=1)) if len(active) > 1 else 0.0
        sd_er = float(np.std(er, ddof=1)) if len(er) > 1 else 0.0
        beta = float(np.cov(returns, rb, ddof=1)[0, 1] / (sd_b ** 2)) if sd_b else None
        excess_path = np.r_[1.0, relative]
        metrics.update(benchmark_return=float(bench[-1] - 1), excess_return=float(relative[-1] - 1),
                       beta=beta, alpha=float(np.mean(returns - rf - beta * (rb - rf)) * trading_days) if beta is not None else None,
                       benchmark_volatility=sd_b * math.sqrt(trading_days),
                       information_ratio=float(np.mean(active) / sd_active * math.sqrt(trading_days)) if sd_active else None,
                       daily_excess_return=float(np.mean(er)),
                       excess_sharpe=float(np.mean(er) / sd_er * math.sqrt(trading_days)) if sd_er else None,
                       excess_max_drawdown=float(-np.min(excess_path / np.maximum.accumulate(excess_path) - 1)))
        daily["benchmark_return"] = bench - 1
        daily["excess_return"] = relative - 1
    prior = initial
    yearly = {}
    for year, group in daily.groupby(daily.index.str[:4]):
        end = float(group.equity.iloc[-1])
        yearly[year] = end / prior - 1
        prior = end
    return metrics, daily.reset_index(), yearly


LABELS = {
    "strategy_return": "策略收益", "annualized_return": "策略年化收益", "excess_return": "超额收益",
    "benchmark_return": "基准收益", "alpha": "阿尔法", "beta": "贝塔", "sharpe": "夏普比率",
    "trade_win_rate": "交易胜率", "profit_loss_ratio": "盈亏比", "max_drawdown": "最大回撤",
    "sortino": "索提诺比率", "daily_excess_return": "日均超额收益", "excess_max_drawdown": "超额最大回撤",
    "excess_sharpe": "超额收益夏普比率", "daily_win_rate": "日胜率", "winning_trades": "盈利次数",
    "losing_trades": "亏损次数", "information_ratio": "信息比率", "strategy_volatility": "策略波动率",
    "benchmark_volatility": "基准波动率"}
PERCENT = {"strategy_return", "annualized_return", "excess_return", "benchmark_return", "max_drawdown",
           "excess_max_drawdown", "daily_excess_return", "trade_win_rate", "daily_win_rate",
           "strategy_volatility", "benchmark_volatility"}


def write_report(out, result, manifest, benchmark=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    import plotly.graph_objects as go
    out = Path(out)
    if out.exists():
        raise FileExistsError("Report directory exists; select a new output")
    out.mkdir(parents=True, mode=0o700)
    closed = episodes(result.book.ledger.entries)
    cashflows = getattr(result, "cashflows", [])
    for trade in closed:
        dividends = sum((Decimal(row["amount"]) for row in cashflows
                         if row["symbol"] == trade["symbol"] and
                         trade["opened"][:10] <= row["record_date"] <= trade["closed"][:10]), Decimal("0"))
        if cashflows:
            trade["dividend_income"] = dividends
            trade["pnl"] += dividends
            trade["return_rate"] = float(trade["pnl"] / trade["buy_cost"])
    metrics, daily, yearly = overview(result.equity, result.book.ledger.initial_cash, closed, benchmark,
                                     manifest.get("trading_days", 252), manifest.get("risk_free", 0.0))
    fills = []
    for entry in result.book.ledger.entries:
        row = asdict(entry.fill)
        row.update(cash_after=entry.cash_after, quantity_after=entry.quantity_after,
                   cost_after=entry.cost_after, realized_pnl=entry.realized_pnl)
        fills.append(json_value(row))
    orders = []
    for order in result.book.orders.values():
        row = asdict(order.request)
        row.update(status=order.status, filled=order.filled, remaining=order.remaining,
                   gross=order.gross, fees=order.fees, reason=order.reason)
        orders.append(json_value(row))
    summary = {"metrics": metrics, "yearly_returns": yearly, "orders": len(orders), "fills": len(fills),
               "data_errors": manifest.get("data_errors", 0),
               "performance_valid": manifest.get("performance_valid", True),
               "dividend_income": float(sum((Decimal(x["amount"]) for x in cashflows), Decimal("0"))),
               "nav_fallback_successes": manifest.get("nav_fallback_successes", 0),
               "unresolved_nav_signal_days": manifest.get("unresolved_nav_signal_days", 0),
               "unavailable": {k: "需要基准数据、完整平仓交易或非零收益波动" for k, v in metrics.items() if v is None},
               "formulas": {"excess_return": "strategy_wealth / benchmark_wealth - 1",
                            "trade": "position from zero to zero, grouped by symbol; weighted average entry cost",
                            "annualization": "252 trading days by default; sample daily volatility ddof=1",
                            "alpha": "annualized mean daily excess return minus beta * benchmark excess return",
                            "sortino": "daily excess mean / RMS of negative daily excess * sqrt(trading_days)",
                            "excess_sharpe": "relative-wealth daily returns, zero risk-free rate",
                            "fees": "minimum commission per order, cumulative across partial fills"}}
    for name, data in (("manifest.json", manifest), ("summary.json", summary)):
        with (out / name).open("x", encoding="utf-8") as stream:
            json.dump(json_value(data), stream, ensure_ascii=False, indent=2, allow_nan=False)
    daily.to_csv(out / "equity.csv", index=False)
    event_rows = [dict(event="run", **json_value(manifest))] + result.events
    event_rows = [dict(row, sequence=i) for i, row in enumerate(event_rows)]
    with (out / "events.jsonl").open("x", encoding="utf-8") as stream:
        for row in event_rows:
            stream.write(json.dumps(json_value(row), ensure_ascii=False, allow_nan=False) + "\n")
    with (out / "backtest.log").open("x", encoding="utf-8") as stream:
        for row in event_rows:
            stream.write(str(row.get("time", row.get("session", ""))) + " " + row["event"] + " " +
                         json.dumps(json_value(row), ensure_ascii=False) + "\n")
    wb = Workbook()
    wb.remove(wb.active)
    sheets = {"委托": orders, "成交": fills, "已平仓交易": json_value(closed),
              "账户": daily.to_dict("records"), "持仓": result.positions,
              "收益概况": [{"指标": LABELS.get(k, k), "值": v} for k, v in metrics.items()]}
    if hasattr(result, "cashflows"):
        sheets["现金分红"] = cashflows
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        keys = list(rows[0]) if rows else ["说明"]
        ws.append(keys)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="24517A")
        for row in rows:
            numeric = {"price", "fee", "cash_after", "cost_after", "realized_pnl", "limit_price",
                       "fee_budget", "gross", "fees", "buy_cost", "pnl", "cash", "equity",
                       "reserved_cash", "cost", "market_value", "dividend_income", "amount"}
            values = [float(row[k]) if k in numeric and row.get(k) is not None else row.get(k) for k in keys]
            # Explicit text cells prevent spreadsheet formula injection.
            ws.append(values)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        if not rows:
            ws.append(["本次运行无此类记录"])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(40, max(14, len(str(col[0].value)) * 2 + 4))
    wb.save(out / "transactions.xlsx")
    fig = go.Figure()
    for key, label in (("strategy_return", "策略收益"), ("benchmark_return", manifest.get("benchmark", "基准收益")),
                       ("excess_return", "超额收益")):
        if key in daily:
            fig.add_trace(go.Scatter(x=daily.date, y=daily[key], name=label))
    fig.update_layout(template="plotly_white", height=520, hovermode="x unified", yaxis_tickformat=".1%",
                      xaxis=dict(rangeslider=dict(visible=True), rangeselector=dict(buttons=[
                          dict(count=1, label="1个月", step="month", stepmode="backward"),
                          dict(count=1, label="1年", step="year", stepmode="backward"), dict(step="all", label="全部")])))
    cards = []
    for key, label in LABELS.items():
        value = metrics[key]
        formatted = "—" if value is None else (f"{value:.2%}" if key in PERCENT else f"{value:.3f}")
        cards.append(f"<div class='card'><small>{label}</small><strong>{formatted}</strong></div>")
    plot = fig.to_html(full_html=False, include_plotlyjs=True)
    error_note = (f"<p style='color:#b00020;font-weight:bold'>ERROR：本次有 {manifest['data_errors']} 条数据错误。成功净值回退 {manifest.get('nav_fallback_successes', 0)} 次；回退后仍缺净值的信号日 {manifest.get('unresolved_nav_signal_days', 0)} 天。请结合错误与回退记录评估收益。</p>"
                  if manifest.get("data_errors") else "")
    page = ("<!doctype html><html lang='zh-CN'><meta charset='utf-8'><title>本地回测报告</title>"
            "<style>body{font-family:system-ui;margin:24px;background:#f6f8fb;color:#243547}"
            ".cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:12px}"
            ".card{background:white;padding:18px;border-radius:8px}small{display:block;color:#65758b}"
            "strong{display:block;font-size:23px;margin-top:10px}</style><h1>本地回测报告</h1><p>" +
            html.escape(manifest.get("description", "")) + "</p>" + error_note + "<div class='cards'>" + "".join(cards) +
            "</div><p>缺失指标显示 —；统计口径见 summary.json。点击图例切换曲线，拖动下方滑块选择日期。</p>" + plot + "</html>")
    with (out / "report.html").open("x", encoding="utf-8") as stream:
        stream.write(page)
    readme = ("# 本地数量回测样例\n\n" + manifest.get("description", "") +
              "\n\n- report.html：离线收益概况与交互曲线。\n- transactions.xlsx：六张账本和统计工作表。\n"
              "- backtest.log / events.jsonl：可读与结构化日志。\n- equity.csv：每日权益及收益。\n"
              "- summary.json：指标、年度收益、缺失原因及公式。\n- manifest.json：运行参数和来源。\n"
              "- checksums.json：输出校验值。\n\n数据频率、撮合、净值可见性和证券规则见 manifest.json；跨平台一致性需要校准。\n")
    with (out / "README.md").open("x", encoding="utf-8") as stream:
        stream.write(readme)
    checks = {file.name: hashlib.sha256(file.read_bytes()).hexdigest() for file in out.iterdir() if file.is_file()}
    with (out / "checksums.json").open("x", encoding="utf-8") as stream:
        json.dump(checks, stream, indent=2)
    return summary
